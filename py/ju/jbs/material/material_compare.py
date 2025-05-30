import csv
import time
import pandas
import xlrd
from toollib.guid import SnowFlake
import openpyxl
from openpyxl.utils import datetime as xl_datetime

from pymysql_comm import UsingOnlineOMS as oms_online
from pymysql_comm import UsingOnlineOMS as oms_dev
import datetime
import json
from typing import List, Tuple


def get_unit():
    sql = "SELECT * FROM oms_product.`unit`"
    with oms_online() as um:
        um.cursor.execute(sql)
        return um.cursor.fetchall()


def get_unit_map():
    units = get_unit()
    unit_name_map = {}
    unit_code_map = {}
    for unit in units:
        unit_name_map[unit['unit_name']] = unit
        unit_code_map[unit['unit_code']] = unit
    return unit_name_map, unit_code_map


def get_material(material_code):
    sql = "SELECT * FROM oms_product.`product_sku` where goods_code = \'{material_code}\'".format(
        material_code=material_code)
    with oms_online() as um:
        um.cursor.execute(sql)
        materials = um.cursor.fetchall()
        return materials[0]

def get_materials(material_code):
    in_sql = ','.join(repr(str(code)) for code in material_code)
    sql = "SELECT * FROM oms_product.`product_sku` where goods_code in ({in_sql})".format(in_sql=in_sql)
    with oms_online() as um:
        um.cursor.execute(sql)
        return um.cursor.fetchall()

def compare_(materials):
    error_material = {}
    unit_name_map, unit_code_map = get_unit_map()

    material_codes = []
    for material in materials:
        material_codes.append(material['code'])

    exist_materials = get_materials(material_codes)
    exist_material_map = {}
    for exist_material in exist_materials:
        exist_material_map[exist_material['goods_code']] = exist_material
        # if exist_material['base_unit_code'] == '':
        #     print(exist_material['goods_code'])

    for material in materials:
        if material['code'] is None:
            continue
        exist_material = exist_material_map[material['code']]
        if exist_material is None:
            error_material[material['code']] = '物料不存在'
            continue

        if exist_material['base_unit_code'] not in unit_code_map:
            continue
        old_unit = unit_code_map[exist_material['base_unit_code']]

        if material['unit'] not in unit_name_map:
            error_material[material['code']] = '单位不存在'
            continue
        new_unit = unit_name_map[material['unit']]

        if old_unit['accuracy'] > new_unit['accuracy']:
            error_material[material['code']] = '新单位精度小于现单位精度'
            continue
    return error_material


def generate_update_sql(materials):
    update_factory_material_sql_format = "update `oms_product`.`factory_material` set base_unit_code = '{unit_code}' where material_code = '{material_code}';"

    update_material_sql_format = "update `oms_product`.`product_sku` set base_unit_code = '{unit_code}' where goods_code  = '{material_code}' and goods_type = 2;"

    file = open("C:\\Users\\pc\\Desktop\\物料单位更新.sql", "w", encoding='utf-8')

    unit_name_map, unit_code_map = get_unit_map()
    material_codes = []
    for material in materials:
        material_codes.append(material['code'])
    exist_materials = get_materials(material_codes)
    exist_material_map = {}
    for exist_material in exist_materials:
        exist_material_map[exist_material['goods_code']] = exist_material
    for material in materials:
        if material['code'] is None:
            continue
        exist_material = exist_material_map[material['code']]
        if exist_material is None:
            print("物料不存在：{code}".format(code = material['code']))
            continue
        update_fm_sql = update_factory_material_sql_format.format(unit_code=unit_name_map[material['unit']]['unit_code'],
                                                material_code=material['code'])
        update_m_sql = update_material_sql_format.format(
            unit_code=unit_name_map[material['unit']]['unit_code'],
            material_code=material['code'])
        file.write(update_fm_sql+'\n')
        file.write(update_m_sql+'\n')
        file.write("\n")
    pass


def read_excel(filename):
    workbook = openpyxl.load_workbook(filename)
    sheet = workbook["Sheet1"]
    error_material = {}
    read_material = []
    for rowIndex in range(1, sheet.max_row + 1):
        material_code = sheet.cell(row=rowIndex, column=1).value
        material_name = sheet.cell(row=rowIndex, column=2).value
        unit_name = sheet.cell(row=rowIndex, column=3).value

        temp = {}
        temp['code'] = material_code
        temp['name'] = material_name
        temp['unit'] = unit_name

        read_material.append(temp)

    compare_result = compare_(materials=read_material)
    error_material.update(compare_result)
    for key, value in error_material.items():
        print(key, value)

    generate_update_sql(read_material)






if __name__ == '__main__':
    read_excel('C:\\Users\\pc\\Desktop\\在用物料单位修改.xlsx')

