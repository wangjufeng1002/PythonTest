# 设备初始化

from toollib.guid import SnowFlake
import openpyxl
from py.ju.jbs.utils.pymysql_comm import UsingOnlineOMS as oms_dev
import datetime
import json

factory_code = {"兴平佳四工厂": "GC0117",
                "兴平佳三工厂": "GC0116",
                "兴平佳二工厂": "GC0115",
                "兴平佳一工厂": "GC0114"}
unit_code = {
    "PCS": "UN1792812390243696641",
    "Pcs": "UN1792812390243696641",
    "个": "UN1792812390243696640"
}
snow = SnowFlake()


def out_put_sql(equipment_list: list, table, file_name):
    sql = '''insert into mes_pms.{table}({field_sql}) values({value_sql});'''
    file = open("D:\\项目相关\\MES2.0\\基础资料\\{file_name}.sql".format(file_name=file_name), "w", encoding='utf-8')

    for equipment in equipment_list:
        field_sql = ''
        value_sql = ''
        for field in equipment:
            if equipment[field] is not None:
                field_sql += '`' + field + '`,'
                if type(equipment[field]) == int or type(equipment[field]) == float:
                    value_sql += str(equipment[field]) + ','
                if type(equipment[field]) == str:
                    value_sql += '\'' + str(equipment[field]) + '\','
                if type(equipment[field]) == dict:
                    value_sql += '\'' + json.dumps(equipment[field], ensure_ascii=False).replace('\"', '\\"') + '\','
                if type(equipment[field]) == bool:
                    value_sql += str(1 if equipment[field] is True else 0) + ','
                if type(equipment[field]) == datetime.datetime:
                    value_sql += '\'' + str(equipment[field].strftime('%Y-%m-%d %H:%M:%S')) + '\','
        field_sql = field_sql[:-1]
        value_sql = value_sql[:-1]
        print(sql.format(table=table, field_sql=field_sql, value_sql=value_sql))
        file.write(sql.format(table=table, field_sql=field_sql, value_sql=value_sql) + '\n')


def read_supplier_sheet(supplier_sheet):
    supplier_map = {}
    for rowIndex in range(2, supplier_sheet.max_row + 1):
        code = supplier_sheet.cell(row=rowIndex, column=1).value
        name = supplier_sheet.cell(row=rowIndex, column=2).value
        supplier_map[name] = code

    return supplier_map


def read_excel(filename):
    workbook = openpyxl.load_workbook(filename)

    supplier_sheet = workbook["供应商"]
    equipment_sheet = workbook['注塑机']
    equipment_help_sheet = workbook["辅助生产设备"]
    equipment_param_sheet = workbook['设备参数基本信息']

    #
    supplier_map = read_supplier_sheet(supplier_sheet)

    # 注塑机
    # equipment_list = read_equipment_sheet(equipment_sheet,supplier_map)
    # out_put_sql(equipment_list, 'equipment', "设备初始化SQL")
    #
    # # 辅助设备
    # help_equipment_list = read_equipment_help_sheet(equipment_help_sheet,supplier_map)
    # out_put_sql(help_equipment_list, 'equipment', "辅助设备初始化SQL")

    param_id_map = get_all_produce_param()
    # 设备生产参数
    equipment_param, equipment_sides, equipment_works = read_equipment_param_sheet(equipment_param_sheet, param_id_map)
    out_put_cycle_update_sql(equipment_param, 'equipment_produce_parameter', "设备参数初始化SQL")
    out_put_sql(equipment_sides, 'equipment_side', "设备周边初始化SQL")
    out_put_sql(equipment_works, 'equipment_worker', "设备人工初始化SQL")


# 注塑机
def read_equipment_sheet(equipment_sheet, supplier_map):
    rows = equipment_sheet.rows
    equipment_list = []

    title_index_map = {}
    # print(equipment_sheet.max_column)

    for cowIndex in range(1, equipment_sheet.max_column + 1):
        title = equipment_sheet.cell(1, cowIndex).value
        title_index_map[title] = cowIndex

    equipment_list = []
    for rowIndex in range(2, equipment_sheet.max_row + 1):
        if equipment_sheet.cell(rowIndex, title_index_map['设备编码']).value is None:
            break
        equipment_field = {}
        equipment_field['equipment_id'] = 'E' + str(snow.gen_uid())
        equipment_field['equipment_code'] = equipment_sheet.cell(rowIndex, title_index_map['设备编码']).value
        equipment_field['equipment_name'] = equipment_sheet.cell(rowIndex, title_index_map['设备名称']).value
        equipment_field['supplier_code'] = supplier_map[equipment_sheet.cell(rowIndex, title_index_map['制造商']).value]
        equipment_field['clamping_device_type'] = get_clamping_type(
            equipment_sheet.cell(rowIndex, title_index_map['合膜装置']).value)
        equipment_field['equipment_type'] = 1
        equipment_field['auth_status'] = 10
        equipment_field['bar_code'] = 1
        equipment_field['equipment_status'] = 10
        equipment_field['factory_code'] = get_factory_code(
            equipment_sheet.cell(rowIndex, title_index_map['所属工厂']).value)
        equipment_field['baking_equipment'] = True if equipment_sheet.cell(rowIndex, title_index_map[
            '烘料设备']).value == '是' else False
        equipment_field['measure_unit_code'] = get_unit_code(
            equipment_sheet.cell(rowIndex, title_index_map['计量单位']).value)
        equipment_field['equipment_model'] = equipment_sheet.cell(rowIndex, title_index_map['规格型号']).value
        equipment_field['bar_code'] = equipment_sheet.cell(rowIndex, title_index_map['设备资产编码']).value
        equipment_field['used_start_date'] = equipment_sheet.cell(rowIndex, title_index_map['开始使用时间']).value
        equipment_field['equipment_tonnage'] = equipment_sheet.cell(rowIndex, title_index_map['机台吨位(T)']).value
        equipment_field['equipment_location'] = equipment_sheet.cell(rowIndex, title_index_map['位置/机台号']).value
        equipment_field['work_hour_max'] = equipment_sheet.cell(rowIndex, title_index_map['最大工作时长(h)']).value
        equipment_field['productivity'] = equipment_sheet.cell(rowIndex, title_index_map['生产效率(只/小时)']).value
        equipment_field['rated_power'] = equipment_sheet.cell(rowIndex, title_index_map['额定功率(kw)']).value
        equipment_field['equipment_cost'] = equipment_sheet.cell(rowIndex, title_index_map['资产原值(元)']).value
        equipment_field['energy_consumption'] = equipment_sheet.cell(rowIndex, title_index_map['实际能耗(%)']).value
        equipment_field['depreciation_period'] = equipment_sheet.cell(rowIndex, title_index_map['折旧年限(年)']).value
        equipment_field['remark'] = equipment_sheet.cell(rowIndex, title_index_map['备注']).value

        equipment_param = {}
        equipment_param['equipmentCode'] = equipment_sheet.cell(rowIndex, title_index_map['设备编码']).value
        equipment_param['screwDiameter'] = equipment_sheet.cell(rowIndex, title_index_map['螺杆直径(mm)']).value
        equipment_param['glueVolume'] = equipment_sheet.cell(rowIndex, title_index_map['射胶量(g)']).value
        equipment_param['takeProductSpace'] = equipment_sheet.cell(rowIndex, title_index_map['取产品空间(mm)']).value
        equipment_param['lockMouldMax'] = equipment_sheet.cell(rowIndex, title_index_map['锁模力最大(ton)']).value
        equipment_param['openMould'] = equipment_sheet.cell(rowIndex, title_index_map['开模力(ton)']).value
        equipment_param['openMouldTravel'] = equipment_sheet.cell(rowIndex, title_index_map['开模行程(mm)']).value
        equipment_param['guidePoleHorizontal'] = equipment_sheet.cell(rowIndex,
                                                                      title_index_map['导柱内距水平(mm)']).value
        equipment_param['guidePoleVertical'] = equipment_sheet.cell(rowIndex, title_index_map['导柱内距垂直(mm)']).value
        equipment_param['mouldCapacityThin'] = equipment_sheet.cell(rowIndex, title_index_map['模容量最薄(mm)']).value
        equipment_param['mouldCapacityThick'] = equipment_sheet.cell(rowIndex, title_index_map['模容量最厚(mm)']).value
        equipment_param['mouldMaxDistance'] = equipment_sheet.cell(rowIndex, title_index_map['模板最大距离(mm)']).value
        equipment_param['oilPumpMaxPower'] = equipment_sheet.cell(rowIndex, title_index_map['油泵电机功率(kw)']).value
        equipment_param['electricityHeat'] = equipment_sheet.cell(rowIndex, title_index_map['电热量(kw)']).value
        equipment_field['equipment_param'] = equipment_param

        equipment_list.append(equipment_field)

    print(len(equipment_list))
    return equipment_list


# 辅助设备
def read_equipment_help_sheet(equipment_sheet, supplier_map):
    title_index_map = {}
    for cowIndex in range(1, equipment_sheet.max_column + 1):
        title = equipment_sheet.cell(1, cowIndex).value
        title_index_map[title] = cowIndex

    equipment_list = []
    for rowIndex in range(2, equipment_sheet.max_row + 1):
        if equipment_sheet.cell(rowIndex, title_index_map['设备编码']).value is None:
            break
        equipment_field = {}
        equipment_field['equipment_id'] = 'E' + str(snow.gen_uid())
        equipment_field['equipment_code'] = equipment_sheet.cell(rowIndex, title_index_map['设备编码']).value
        equipment_field['equipment_name'] = equipment_sheet.cell(rowIndex, title_index_map['设备名称']).value
        equipment_field['supplier_code'] = supplier_map[equipment_sheet.cell(rowIndex, title_index_map['制造商']).value]
        equipment_field['equipment_type'] = 2
        equipment_field['auth_status'] = 10
        equipment_field['bar_code'] = 1
        equipment_field['equipment_status'] = 10
        equipment_field['factory_code'] = get_factory_code(
            equipment_sheet.cell(rowIndex, title_index_map['所属工厂']).value)
        equipment_field['baking_equipment'] = True if equipment_sheet.cell(rowIndex, title_index_map[
            '烘料设备']).value == '是' else False
        equipment_field['measure_unit_code'] = get_unit_code(
            equipment_sheet.cell(rowIndex, title_index_map['计量单位']).value)
        equipment_field['equipment_model'] = equipment_sheet.cell(rowIndex, title_index_map['规格型号']).value
        equipment_field['bar_code'] = equipment_sheet.cell(rowIndex, title_index_map['设备资产编码']).value
        equipment_field['used_start_date'] = equipment_sheet.cell(rowIndex, title_index_map['开始使用时间']).value
        equipment_field['equipment_tonnage'] = equipment_sheet.cell(rowIndex, title_index_map['机台吨位(T)']).value
        equipment_field['equipment_location'] = equipment_sheet.cell(rowIndex, title_index_map['位置/机台号']).value
        equipment_field['work_hour_max'] = equipment_sheet.cell(rowIndex, title_index_map['最大工作时长(h)']).value
        equipment_field['productivity'] = equipment_sheet.cell(rowIndex, title_index_map['生产效率(只/小时)']).value
        equipment_field['rated_power'] = equipment_sheet.cell(rowIndex, title_index_map['额定功率(kw)']).value
        equipment_field['equipment_cost'] = equipment_sheet.cell(rowIndex, title_index_map['资产原值(元)']).value
        equipment_field['energy_consumption'] = equipment_sheet.cell(rowIndex, title_index_map['实际能耗(%)']).value
        equipment_field['depreciation_period'] = equipment_sheet.cell(rowIndex, title_index_map['折旧年限(年)']).value
        equipment_field['remark'] = equipment_sheet.cell(rowIndex, title_index_map['备注']).value
        equipment_field['equipment_param'] = '{}'
        if equipment_field['supplier_code'] is None:
            equipment_field['supplier_code'] = '未知'
        if equipment_field['equipment_tonnage'] is not None:
            tonnage = equipment_field['equipment_tonnage']
            if 'T' in str(tonnage):
                equipment_field['equipment_tonnage'] = float(str(tonnage).replace('T', ''))
        equipment_list.append(equipment_field)
    print(len(equipment_list))
    return equipment_list


# 设备生产参数
def read_equipment_param_sheet(equipment_sheet, param_id_map):
    title_index_map = {}
    for cowIndex in range(1, equipment_sheet.max_column + 1):
        title = equipment_sheet.cell(1, cowIndex).value
        title_index_map[title] = cowIndex

    equipment_list = []
    equipment_worker_list = []
    equipment_side_list = []
    for rowIndex in range(2, equipment_sheet.max_row + 1):
        if equipment_sheet.cell(rowIndex, title_index_map['设备编码']).value is None:
            break
        if equipment_sheet.cell(rowIndex, title_index_map['模具编号']).value is None:
            continue
        equipment_field = {}

        equipment_code = equipment_sheet.cell(rowIndex, title_index_map['设备编码']).value
        mould_code = equipment_sheet.cell(rowIndex, title_index_map['模具编号']).value
        if equipment_code + mould_code not in param_id_map:
            continue
        else:
            equipment_field['produce_param_id'] = param_id_map[equipment_code + mould_code]

        #equipment_field['produce_param_id'] = 'EP' + str(snow.gen_uid())
        equipment_field['equipment_code'] = equipment_sheet.cell(rowIndex, title_index_map['设备编码']).value
        equipment_field['connect_status'] = 0
        equipment_field['produce_status'] = 10
        equipment_field['operate_type'] = 10
        equipment_field['mould_code'] = equipment_sheet.cell(rowIndex, title_index_map['模具编号']).value
        equipment_field['mould_serial_no'] = ''
        equipment_field['priority'] = 1
        equipment_field['process_param'] = {
            "cycleTime": equipment_sheet.cell(rowIndex, title_index_map['循环时间(s)']).value}
        equipment_field['create_by'] = '963903739941949541'
        equipment_list.append(equipment_field)

        side_code = equipment_sheet.cell(rowIndex, title_index_map['机械手编码']).value
        if side_code is not None and side_code != '' :
            equipment_side_1 = {}
            equipment_side_1['equipment_side_id'] = 'EPS' + str(snow.gen_uid())
            equipment_side_1['equipment_produce_param_id'] = equipment_field['produce_param_id']
            equipment_side_1['equipment_side_code'] = side_code
            equipment_side_list.append(equipment_side_1)
        if equipment_sheet.cell(rowIndex, title_index_map['烘料桶']).value is not None:
            equipment_side_2 = {}
            equipment_side_2['equipment_side_id'] = 'EPS' + str(snow.gen_uid())
            equipment_side_2['equipment_produce_param_id'] = equipment_field['produce_param_id']
            equipment_side_2['equipment_side_code'] = equipment_sheet.cell(rowIndex, title_index_map[
                '烘料桶']).value
            equipment_side_list.append(equipment_side_2)

        worker = equipment_sheet.cell(rowIndex, title_index_map['调剂师人数']).value
        if worker is not None and worker != '' :
            equipment_worker = {}
            equipment_worker['equipment_produce_param_id'] = equipment_field['produce_param_id']
            equipment_worker['pilot_num'] = equipment_sheet.cell(rowIndex, title_index_map['调剂师人数']).value
            equipment_worker['injection_num'] = equipment_sheet.cell(rowIndex, title_index_map['注塑操作工人数']).value
            equipment_worker['stir_num'] = equipment_sheet.cell(rowIndex, title_index_map['拌料工人数']).value
            equipment_worker['crush_num'] = equipment_sheet.cell(rowIndex, title_index_map['碎料工人数']).value
            equipment_worker_list.append(equipment_worker)

    unique = {}
    for e_p in equipment_list:
        if e_p['equipment_code'] + e_p['mould_code'] not in unique:
            unique[e_p['equipment_code'] + e_p['mould_code']] = 1
        else:
            unique[e_p['equipment_code'] + e_p['mould_code']] = unique[e_p['equipment_code'] + e_p['mould_code']] + 1
            print(e_p['equipment_code'] + "<=====> " + e_p['mould_code'])
    return equipment_list, equipment_side_list, equipment_worker_list


def out_put_cycle_update_sql(equipment_param_list, tale_name, file_name):
    sql = '''UPDATE equipment_produce_parameter SET process_param = {process_param} WHERE produce_param_id = {produce_param_id};'''
    file = open("D:\\项目相关\\MES2.0\\基础资料\\{file_name}.sql".format(file_name=file_name), "w", encoding='utf-8')


    for param in equipment_param_list:
        produce_param_id = '\'' + str(param['produce_param_id']) + '\''
        process_param = '\'' + json.dumps(param['process_param'], ensure_ascii=False).replace('\"','\\"') + '\''
        file.write(sql.format(table=tale_name, produce_param_id=produce_param_id,
                              process_param=process_param) +  '\n')


def get_value(sheet, row, col):
    return sheet.cell(row, col).value


def get_factory_code(factory_name):
    return factory_code[factory_name]


def get_unit_code(unit_name):
    return unit_code[unit_name]


def parse_date(date, format):
    return datetime.datetime.strptime(date, format)


def format_date(date, format):
    return date.strftime(format)


def get_clamping_type(name):
    if name == '二板式':
        return 1
    if name == '三板式':
        return 2
    if name == '四板式':
        return 3
    if name == '五板式':
        return 4


def get_produce_param(page_num, page_size):
    page_from = (page_num - 1) * page_size
    sql = "SELECT * FROM mes_pms.`equipment_produce_parameter` limit %d,%d" % (
        page_from, page_size)
    with oms_dev() as um:
        um.cursor.execute(sql)
        return um.cursor.fetchall()


def get_all_produce_param():
    page_num = 1
    page_size = 100
    param_map = {}
    while True:
        param_res = get_produce_param(page_num, page_size)
        if len(param_res) == 0:
            break
        for param in param_res:
            param_map[param['equipment_code'] + param['mould_code']] = param["produce_param_id"]
        page_num += 1
    return param_map


if __name__ == '__main__':
    snowflake = SnowFlake(worker_id=1, datacenter_id=1)
    print(snowflake.gen_uid())

    read_excel("D:\\项目相关\\MES2.0\\基础资料\\设备-基础信息.xlsx")
