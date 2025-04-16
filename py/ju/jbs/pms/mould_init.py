# 设备初始化

import csv
import time
import pandas
import xlrd
from toollib.guid import SnowFlake
import openpyxl
from openpyxl.utils import datetime as xl_datetime
from pymysql_comm import UsingOnlineOMS as oms_online
from pymysql_comm import UsingTest as oms_dev
import datetime
import json

snow = SnowFlake()
def equipment_mould():
    sql = "INSERT INTO `mes_pms`.`mould_equipment`(`mould_equipment_id`, `mould_code`, `equipment_code`, `priority`) VALUES ('{id}', '{m_code}', '{e_code}', '{p}');";
    filename = 'D:\\项目相关\\MES2.0\\基础资料\\模具与注塑机绑定关系.xlsx'
    workbook = openpyxl.load_workbook(filename)
    supplier_sheet = workbook["Sheet1"]
    file = open("D:\\项目相关\\MES2.0\\基础资料\\{file_name}.sql".format(file_name="模具与注塑机绑定关系sql"), "w",
                encoding='utf-8')
    for rowIndex in range(2, supplier_sheet.max_row + 1):
        m_code = supplier_sheet.cell(rowIndex, 1).value
        e_code = supplier_sheet.cell(rowIndex, 2).value
        s = supplier_sheet.cell(rowIndex, 3).value
        sql_format = sql.format(id=snow.gen_uid(), m_code=m_code, e_code=e_code, p=s)
        print(sql_format)
        file.write(sql_format + '\n')
if __name__ == '__main__':


    print("")