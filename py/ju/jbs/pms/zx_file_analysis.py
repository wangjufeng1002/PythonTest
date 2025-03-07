import openpyxl
import json
if __name__ == '__main__':
    workbook = openpyxl.load_workbook("D:\\项目相关\\MES2.0\\注塑机文档\\震雄\\解析规则\\处理\\佳帮手-MPC7-8605-test.xlsx")
    sheet_ = workbook['Sheet1']

    nodes =[]
    for rowIndex in range(10, 459):
        if sheet_.cell(row=rowIndex, column=1).value is None:
            continue
        node = {}
        node['pos'] = sheet_.cell(rowIndex, 2).value
        node['size'] = sheet_.cell(rowIndex, 3).value
        desc1 = sheet_.cell(rowIndex, 4).value
        desc2 = sheet_.cell(rowIndex, 10).value
        node['name'] = str(desc2) + str(desc1)
        node['type'] = 'UINT16'
        node['precision'] = sheet_.cell(rowIndex, 6).value
        node['jsonPath'] = '$.'
        nodes.append(node)

    print(json.dumps(nodes, ensure_ascii=False))

