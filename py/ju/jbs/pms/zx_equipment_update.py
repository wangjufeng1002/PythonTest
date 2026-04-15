import openpyxl

ip_1 = "124.115.116.142"
ip_2 = "36.163.199.14"
ip_3 = "61.185.16.90"

update_sql = "update `mes_pms`.`equipment` set `ip` = '{ip}',`port`= '{port}' ,`auth_status` = 1 ,`equipment_location` = '{location}',`ctrl_code`='{ctrl_type}',`fuselage_code` = '{fuselage_code}' where `equipment_code` = '{equipment_code}' ;"

def get_ctrl_type(ctrl_type):
    if ctrl_type == 'CPC-6.6':
        return 'CPC6'
    elif ctrl_type == 'CPC-6.0':
        return 'CPC6'
    elif ctrl_type == 'MPC-7.0':
        return 'MPC7'
    elif ctrl_type == 'ai-12':
        return 'MPC7'
    elif ctrl_type == 'ai-02':
        return 'MPC7'
    elif ctrl_type == 'MCP3000' or ctrl_type == 'MCP 3000' :
        return 'MAIMI'
    elif ctrl_type == 'YAQI':
        return 'YAQI'
    return ctrl_type


def read_sheet(sheet, ip):
    title_index_map = {}
    sqls =[]
    for cowIndex in range(1, sheet.max_column + 1):
        title = sheet.cell(3, cowIndex).value
        title_index_map[title] = cowIndex
    for rowIndex in range(4, sheet.max_row + 1):
        equipment_code = sheet.cell(rowIndex, title_index_map['设备编号']).value
        if equipment_code is None:
            continue
        location_code = sheet.cell(rowIndex, title_index_map['机台号']).value
        ctrl_type = sheet.cell(rowIndex, title_index_map['电脑型号']).value
        fuselage_code = sheet.cell(rowIndex, title_index_map['机身编号']).value
        port = 20000 + location_code
        sql = update_sql.format(equipment_code=equipment_code, ip=ip, port=port, location=location_code,
                                fuselage_code=fuselage_code, ctrl_type=get_ctrl_type(ctrl_type))
        sqls.append(sql)
    file = open("D:\\测试文件\\pms\\{file_name}.sql".format(file_name="震雄注塑机授权更新"), "a+", encoding='utf-8')
    for sql in sqls:
        file.write(sql+ "\n")
        print(sql)

def read_excel():
    workbook = openpyxl.load_workbook("D:\项目相关\MES2.0\注塑机文档\震雄\佳帮手震雄注塑机统计-2025-07-16-1-整理.xlsx")
    jia_one_sheet = workbook["佳一"]
    jia_two_sheet = workbook["佳二"]
    jia_three_sheet = workbook["佳三"]

    read_sheet(jia_one_sheet, ip_1)
    read_sheet(jia_two_sheet, ip_2)
    read_sheet(jia_three_sheet, ip_3)

if __name__ == '__main__':
    read_excel()
